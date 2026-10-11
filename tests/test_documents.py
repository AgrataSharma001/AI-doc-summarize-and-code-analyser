import io
import json
import zipfile

import httpx
import pytest
from document_fixtures import docx_bytes, pdf_bytes
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import select

from backend.config import Settings
from backend.documents import extract_document
from backend.main import create_app
from backend.models import Conversation, SourceFile, SourceSpan
from backend.ollama import OllamaAdapter
from backend.sources import SPAN_CHARS, select_evidence

ORIGIN = {"Origin": "http://testserver"}



def source():
    return SourceFile(id='source-id', filename='report')


def test_pdf_pages_chunks_retrieval_and_empty_page_warning():
    file = source()
    warnings = extract_document(file, pdf_bytes('Alpha findings.', None, 'Zebra budget is 240.'), '.pdf')
    assert [span.location for span in file.spans] == ['page 1, lines 1-1', 'page 3, lines 1-1']
    assert [span.position for span in file.spans] == [0, 1]
    assert len({span.id for span in file.spans}) == 2
    assert all(span.file_id == file.id for span in file.spans)
    assert '1 PDF page(s)' in warnings[0]
    assert 'OCR' in warnings[0]
    selected, partial = select_evidence([file], 'zebra budget', 220)
    assert partial and len(selected) == 1
    assert 'Zebra budget' in selected[0]['content']
    assert selected[0]['location'].startswith('page 3,')


def test_docx_body_order_heading_table_and_long_paragraph():
    document = Document()
    document.add_heading('Budget', level=1)
    document.add_paragraph('Before table.')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Item'
    table.cell(0, 1).text = '240 rupees'
    document.add_paragraph('After table.' + 'x' * (SPAN_CHARS * 2))
    result = io.BytesIO()
    document.save(result)
    file = source()
    extract_document(file, result.getvalue(), '.docx')
    texts = [span.content for span in file.spans]
    assert texts[:4] == ['Budget', 'Before table.', 'Item', '240 rupees']
    assert texts[4].startswith('After table.')
    assert all('section 1 (Budget)' in span.location for span in file.spans)
    assert 'paragraph 2' in file.spans[1].location
    assert 'table 1 row 1 cell 2' in file.spans[3].location
    assert all(len(span.content) <= SPAN_CHARS and len(span.location) <= 100 for span in file.spans)
    assert [span.position for span in file.spans] == list(range(len(file.spans)))


@pytest.mark.parametrize('extension,content,expected', [
    ('.pdf', pdf_bytes(None), 'No extractable'),
    ('.pdf', pdf_bytes('Secret', encrypted=True), 'encrypted'),
    ('.pdf', pdf_bytes(), 'no pages'),
    ('.pdf', b'%PDF-1.7\ninvalid', 'corrupt'),
    ('.docx', docx_bytes(''), 'No extractable'),
    ('.docx', b'not a zip', 'corrupt'),
], ids=['blank-pdf', 'encrypted-pdf', 'empty-pdf', 'corrupt-pdf', 'empty-docx', 'corrupt-docx'])
def test_invalid_documents_have_actionable_errors_and_no_partial_spans(extension, content, expected):
    file = source()
    with pytest.raises(HTTPException) as exc:
        extract_document(file, content, extension)
    assert exc.value.status_code == 422
    assert expected in exc.value.detail
    assert file.spans == []


@pytest.mark.parametrize('extension', ['.pdf', '.docx'])
def test_extracted_character_limit_is_atomic(extension, monkeypatch):
    monkeypatch.setattr('backend.documents.MAX_SOURCE_CHARS', 40)
    content = pdf_bytes('x' * 30, 'y' * 30) if extension == '.pdf' else docx_bytes('x' * 60)
    file = source()
    with pytest.raises(HTTPException) as exc:
        extract_document(file, content, extension)
    assert exc.value.status_code == 413
    assert file.spans == []


@pytest.mark.parametrize('limit,value', [('MAX_PDF_PAGES', 1), ('MAX_PDF_PAGE_STREAM', 10),
                                        ('MAX_EXPANDED_BYTES', 10)])
def test_pdf_structure_and_stream_limits(limit, value, monkeypatch):
    monkeypatch.setattr(f'backend.documents.{limit}', value)
    with pytest.raises(HTTPException) as exc:
        extract_document(source(), pdf_bytes('Alpha', 'Beta'), '.pdf')
    assert exc.value.status_code == 413


def test_docx_expansion_limit(monkeypatch):
    monkeypatch.setattr('backend.documents.MAX_EXPANDED_BYTES', 100)
    with pytest.raises(HTTPException) as exc:
        extract_document(source(), docx_bytes(), '.docx')
    assert exc.value.status_code == 413


def test_compressed_pdf_is_limited_during_decompression(monkeypatch):
    monkeypatch.setattr('backend.documents.MAX_PDF_PAGE_STREAM', 100)
    with pytest.raises(HTTPException) as exc:
        extract_document(source(), pdf_bytes('x' * 5000, compressed=True), '.pdf')
    assert exc.value.status_code == 413
    assert 'parser resource limits' in exc.value.detail


def test_docx_merged_cells_are_not_repeated():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 1)).text = 'Merged budget 240'
    result = io.BytesIO()
    document.save(result)
    file = source()
    extract_document(file, result.getvalue(), '.docx')
    assert [span.content for span in file.spans] == ['Merged budget 240']


def test_docx_rejects_dtd_in_auxiliary_part():
    original = io.BytesIO(docx_bytes())
    changed = io.BytesIO()
    with zipfile.ZipFile(original) as archive, zipfile.ZipFile(changed, 'w') as target:
        for entry in archive.infolist():
            data = archive.read(entry)
            if entry.filename == 'word/styles.xml':
                data = b'<!DOCTYPE x [<!ENTITY leak SYSTEM "file:///private">]>' + data
            target.writestr(entry, data)
    with pytest.raises(HTTPException) as exc:
        extract_document(source(), changed.getvalue(), '.docx')
    assert exc.value.status_code == 422
    assert 'private' not in exc.value.detail


@pytest.fixture
def client(tmp_path):
    app = create_app(Settings(database_url=f"sqlite:///{(tmp_path / 'documents.db').as_posix()}"))

    def model(request):
        prompt = json.loads(json.loads(request.content)['messages'][1]['content'])
        return httpx.Response(200, json={'done': True, 'message': {'content': json.dumps({
            'reply': 'A supported answer.', 'source_ids': [s['id'] for s in prompt['source_excerpts']],
        })}})

    app.state.model_adapter = OllamaAdapter(app.state.settings, httpx.MockTransport(model))
    with TestClient(app) as client:
        assert client.post('/api/auth/signup', headers=ORIGIN, json={
            'name': 'Alex', 'email': 'alex@example.com', 'password': 'strong-pass-123',
        }).status_code == 200
        yield client


def ask(client, files=()):
    parts = [(key, (None, value)) for key, value in {
        'mode': 'document', 'message': 'Summarize', 'conversation_id': 'doc-chat', 'history': '[]',
    }.items()]
    return client.post('/api/chat', headers=ORIGIN, files=parts + [('files', item) for item in files])


@pytest.mark.parametrize('extension,content,location', [
    ('.pdf', pdf_bytes('Budget is 240.', None), 'page 1,'),
    ('.docx', docx_bytes('Budget is 240.'), 'paragraph 1,'),
], ids=['pdf', 'docx'])
def test_document_api_persists_sources_and_followup_citations(client, extension, content, location):
    first = ask(client, [(f'report{extension}', content, 'application/octet-stream')])
    assert first.status_code == 200
    answer = first.json()
    assert location in answer['reply']
    if extension == '.pdf':
        assert answer['warnings'] and 'OCR' in answer['reply']
    followup = ask(client)
    assert followup.status_code == 200
    assert followup.json()['source_ids'] == answer['source_ids']
    with client.app.state.session_factory() as db:
        file = db.get(Conversation, 'doc-chat').files[0]
        assert file.spans[0].location.startswith(location)
        assert '240' in file.spans[0].content


def test_invalid_document_batch_never_calls_model_or_saves_sources(client):
    class UnexpectedModel:
        async def answer(self, *args):
            pytest.fail('Invalid extraction must not call the model')

    client.app.state.model_adapter = UnexpectedModel()
    response = ask(client, [('good.docx', docx_bytes(), 'application/octet-stream'),
                            ('scan.pdf', pdf_bytes(None), 'application/pdf')])
    assert response.status_code == 422
    assert 'OCR' in response.json()['error']['message']
    with client.app.state.session_factory() as db:
        assert db.get(Conversation, 'doc-chat') is None
        assert db.scalar(select(SourceSpan)) is None

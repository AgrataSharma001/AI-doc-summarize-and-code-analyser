"""Synthetic in-memory documents, also usable for the live browser test."""

import io
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_bytes(*pages, encrypted=False, compressed=False):
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        if text is not None:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                     NameObject('/Subtype'): NameObject('/Type1'),
                                     NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
                DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
            stream = DecodedStreamObject()
            escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            stream.set_data(f'BT /F1 12 Tf 50 700 Td ({escaped}) Tj ET'.encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
            if compressed:
                page.compress_content_streams()
    if encrypted:
        writer.encrypt('test-password')
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def docx_bytes(text='A short report.'):
    document = Document()
    if text:
        document.add_paragraph(text)
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()



if __name__ == "__main__":
    target = Path(__file__).resolve().parents[1] / ".test-tools" / "fixtures"
    target.mkdir(parents=True, exist_ok=True)
    (target / "orion.pdf").write_bytes(pdf_bytes(
        "Project Orion has a budget of 240 rupees.",
        "The delivery deadline is 18 November 2026.",
        "Mira owns the project.",
    ))
    document = Document()
    document.add_heading("Project Orion", level=1)
    document.add_paragraph("Project Orion has a budget of 240 rupees.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Project owner"
    table.cell(0, 1).text = "Mira"
    document.add_paragraph("The delivery deadline is 18 November 2026.")
    document.save(target / "orion.docx")
    print("Wrote synthetic PDF and DOCX fixtures to .test-tools/fixtures.")

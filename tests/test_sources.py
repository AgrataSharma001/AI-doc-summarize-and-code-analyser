"""Structural extraction and known-answer retrieval checks; no model needed."""

import pytest

from backend.models import SourceFile
from backend.sources import SPAN_CHARS, extract_text, select_evidence


def source_for(text, structure="markdown", filename="notes.md"):
    source = SourceFile(id=filename, filename=filename)
    extract_text(source, text, structure=structure)
    return source


def test_paragraph_boundaries_and_crlf_locations():
    text = "First paragraph.\r\nContinued.\r\n\r\nSecond paragraph.\r\n"
    source = source_for(text, "paragraphs")
    assert [s.location for s in source.spans] == ["lines 1-3", "lines 4-4"]
    assert "".join(s.content for s in source.spans) == text


def test_headings_without_blank_lines_and_section_inheritance():
    text = "# Overview\nIntro.\n## Budget ###\nCosts 42 units.\n\nPayment is due Friday.\n"
    source = source_for(text)
    assert [s.location for s in source.spans] == [
        "section Overview, lines 1-2", "section Budget, lines 3-5",
        "section Budget, lines 6-6",
    ]
    assert "".join(s.content for s in source.spans) == text


def test_setext_headings():
    source = source_for("Overview\n========\nIntro.\nBudget\n------\nCosts 42 units.")
    assert [s.location for s in source.spans] == [
        "section Overview, lines 1-3", "section Budget, lines 4-6",
    ]


@pytest.mark.parametrize("fence", ["```", "~~~~"])
def test_fenced_headings_and_blank_lines_are_literal(fence):
    text = f"# Example\n{fence}python\n# literal heading\n\nLiteral title\n---\n{fence}\nAfter.\n"
    source = source_for(text)
    assert len(source.spans) == 1
    assert source.spans[0].content == text
    assert source.spans[0].location == "section Example, lines 1-8"


def test_long_paragraphs_keep_section_and_line_locations():
    text = "# " + "Title" * 50 + "\n" + "x" * (SPAN_CHARS * 3) + "\n"
    source = source_for(text)
    assert "".join(s.content for s in source.spans) == text
    assert all(len(s.content) <= SPAN_CHARS for s in source.spans)
    assert all(len(s.location) <= 100 for s in source.spans)
    assert all(s.location.startswith("section Title") for s in source.spans)
    assert [s.position for s in source.spans] == list(range(len(source.spans)))
    assert len({s.id for s in source.spans}) == len(source.spans)


def test_python_comments_keep_line_only_locations():
    source = source_for("# A comment\nx = 1\n\n# Another\ny = 2", "lines", "code.py")
    assert len(source.spans) == 1
    assert source.spans[0].location == "lines 1-5"


@pytest.mark.parametrize("position", [0, 9, 19], ids=["early", "middle", "late"])
@pytest.mark.parametrize("structure", ["paragraphs", "markdown"])
def test_known_answer_retrieval_from_long_sources(position, structure):
    paragraphs = [f"Routine record {i}: ordinary administrative information." for i in range(20)]
    paragraphs[position] = "The telescope aperture is 73 centimetres."
    source = source_for("\n\n".join(paragraphs), structure)
    selected, partial = select_evidence([source], "telescope aperture", 300)
    assert partial
    assert len(selected) == 1
    assert "73 centimetres" in selected[0]["content"]
    original = next(s for s in source.spans if "73 centimetres" in s.content)
    assert selected[0]["id"] == original.id
    assert selected[0]["location"] == original.location


def test_retrieval_uses_inherited_markdown_heading():
    source = source_for("# Overview\nRoutine notes.\n# Telescope\n\nAperture: 73 centimetres.")
    selected, partial = select_evidence([source], "telescope aperture", 240)
    assert partial
    assert len(selected) == 1
    assert "73 centimetres" in selected[0]["content"]


def test_multiple_files_order_and_budget():
    first = source_for("Routine text.\n\nThe telescope aperture is 73 centimetres.", filename="one.md")
    second = source_for("Routine text.\n\nThe telescope mass is 42 kilograms.", filename="two.md")
    selected, partial = select_evidence([first, second], "telescope aperture mass", 460)
    assert partial
    assert [s["file_id"] for s in selected] == ["one.md", "two.md"]
    assert all("telescope" in s["content"] for s in selected)
    assert sum(len(s["content"]) + len(s["filename"]) + 160 for s in selected) <= 460
    empty, partial = select_evidence([first], "aperture", 1)
    assert empty == [] and partial


def test_retrieval_unicode_case_and_full_coverage():
    source = source_for("Routine.\n\nCAFÉ opening time is 09:00.")
    selected, partial = select_evidence([source], "café opening", 240)
    assert partial and "09:00" in selected[0]["content"]
    selected, partial = select_evidence([source], "opening", 2000)
    assert not partial
    assert [s["id"] for s in selected] == [s.id for s in source.spans]


@pytest.mark.parametrize("extension", [".pdf", ".docx"])
def test_known_answer_retrieval_keeps_document_locations(extension):
    import io

    from document_fixtures import pdf_bytes
    from docx import Document

    from backend.documents import extract_document

    source = SourceFile(id="report", filename="report" + extension)
    if extension == ".pdf":
        pages = [f"Routine record {i}." for i in range(20)]
        pages[14] = "The telescope aperture is 73 centimetres."
        content = pdf_bytes(*pages)
        expected = "page 15, lines 1-1"
    else:
        document = Document()
        for i in range(20):
            if i == 14:
                document.add_heading("Telescope", level=1)
                document.add_paragraph("Aperture: 73 centimetres.")
            else:
                document.add_paragraph(f"Routine record {i}.")
        output = io.BytesIO()
        document.save(output)
        content = output.getvalue()
        expected = "section 1 (Telescope), paragraph 16, lines 1-1"
    extract_document(source, content, extension)
    selected, partial = select_evidence([source], "telescope aperture", 240)
    assert partial and len(selected) == 1
    assert "73 centimetres" in selected[0]["content"]
    assert selected[0]["location"] == expected
    assert selected[0]["id"] in {s.id for s in source.spans}
